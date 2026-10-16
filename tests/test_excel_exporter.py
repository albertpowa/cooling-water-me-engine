"""Excel exporter tests: structure, live formulas, and Python ↔ Excel parity."""

from pathlib import Path

import openpyxl
import pytest

from src.engine import compare_scenarios, load_config
from src.excel.exporter import SHEETS, build_book, export_workbook

REQUIRED = ["Dashboard", "Process Flow Mass Balance", "Cooling Tower & Corrugated Fill",
            "Pumping & Fan Hydraulics", "Disinfection & Microbial Risk", "Financial & LCC Analysis"]


@pytest.fixture(scope="module")
def workbook_path(tmp_path_factory) -> Path:
    return export_workbook(load_config(), tmp_path_factory.mktemp("xl") / "engine.xlsx")


def test_required_sheets_charts_and_formulas(workbook_path):
    wb = openpyxl.load_workbook(workbook_path)
    assert set(REQUIRED) <= set(wb.sheetnames)
    n_formulas = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
                     if isinstance(c.value, str) and c.value.startswith("="))
    assert n_formulas > 2000
    assert sum(len(ws._charts) for ws in wb.worksheets) >= 6
    text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)
    for fn in ("SUM(", "IF(", "LOG(", "ROUNDUP(", "NPV(", "IRR(", "INDEX(", "MATCH("):
        assert fn in text, fn


def test_inputs_and_outputs_are_color_coded(workbook_path):
    ws = openpyxl.load_workbook(workbook_path)["Dashboard"]
    fills = {ws["C8"].fill.fgColor.rgb[-6:], ws["C30"].fill.fgColor.rgb[-6:]}
    assert "FFF2CC" in fills          # IT load input = soft yellow


def test_no_unresolved_placeholders(workbook_path):
    wb = openpyxl.load_workbook(workbook_path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert "{" not in c.value and c.value != "=0", (ws.title, c.coordinate)


def test_excel_formulas_reproduce_engine(tmp_path):
    """Recalculate every formula with the `formulas` package and compare with Python (slow, ~40 s)."""
    formulas = pytest.importorskip("formulas")
    cfg = load_config()
    book = build_book(cfg)
    path = tmp_path / "parity.xlsx"
    book.wb.save(path)
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()

    def xl(name: str) -> float:
        sheet, cell = book.refs[name].split("!")
        return sol[f"'[{path.name}]{sheet.strip(chr(39)).upper()}'!{cell.replace('$', '')}"].value[0, 0]

    c = compare_scenarios(cfg)
    pairs = [("npv", c.finance.npv), ("capex_total", c.capex["total"]), ("irr", c.finance.irr)]
    for s, r in (("b", c.baseline), ("o", c.optimized)):
        pairs += [(f"intake_{s}", r.intake_gpd), (f"coc_{s}", r.tower.coc), (f"tcold_{s}", r.fill_perf.cold_water_f),
                  (f"kw_total_{s}", r.power_kw["total"]), (f"cost_total_{s}", r.costs["total"]),
                  (f"lsi_hot_{s}", r.scaling["recirc_hot"].lsi), (f"leg_py_{s}", r.risks["legionella"].py),
                  (f"ct_{s}", r.ct["ct"]), (f"leed_pts_{s}", r.leed_points)]
    for name, expected in pairs:
        assert float(xl(name)) == pytest.approx(expected, rel=1e-3, abs=1e-9), name


def test_sheet_registry_matches_prompt_names():
    assert all(name in SHEETS.values() for name in REQUIRED)
