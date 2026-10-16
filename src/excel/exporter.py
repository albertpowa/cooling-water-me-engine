"""openpyxl generator for ``Data_Center_Cooling_ME_Engine.xlsx``.

Every engineering result in the workbook is a live Excel formula — inputs
(soft-yellow, dark-blue text) can be edited in Excel and the whole model
recalculates: mass balance, Merkel fill solver, CoC limits, LSI/RSI, CT,
QMRA, pumps, fans and the 10-year lifecycle cost (NPV / IRR / BEP / SPBT).

Formulas are authored as templates with ``{name}`` placeholders that are
resolved to absolute cell references after every sheet is laid out, so
sheets may reference each other in any order. Two-scenario rows use ``$`` in
a placeholder (``{coc$}``) which expands to ``_b`` (Baseline) / ``_o``
(Optimized).
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence, Tuple, Union

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from ..engine.constants import PATHOGEN_LABELS, PATHOGEN_UNITS, SPECIES, SPECIES_LABELS, SPECIES_UNITS
from ..engine.microbiology_qmra import CT_GIARDIA_3LOG, CT_VIRUS_4LOG
from ..engine.scenario import CHLORINE_CREDIT_KEY, TREATMENT_ORDER

DEFAULT_FILENAME = "Data_Center_Cooling_ME_Engine.xlsx"

# --- Styles -------------------------------------------------------------------------
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")      # soft yellow
OUTPUT_FILL = PatternFill("solid", fgColor="E2EFDA")     # soft green
HEADER_FILL = PatternFill("solid", fgColor="1F3864")     # dark blue
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
CHECK_OK_FILL = PatternFill("solid", fgColor="C6EFCE")
CHECK_BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
INPUT_FONT = Font(name="Calibri", size=11, color="1F3864")
OUTPUT_FONT = Font(name="Calibri", size=11, color="000000")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=18, bold=True, color="1F3864")
SUB_FONT = Font(name="Calibri", size=11, italic=True, color="595959")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="1F3864")
NOTE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

F_INT = "#,##0"
F_1 = "#,##0.0"
F_2 = "#,##0.00"
F_3 = "0.000"
F_4 = "0.0000"
F_SCI = "0.00E+00"
F_USD = "$#,##0"
F_PCT = "0.0%"
F_PCT3 = "0.000%"

SHEETS = {
    "dash": "Dashboard",
    "mb": "Process Flow Mass Balance",
    "ct": "Cooling Tower & Corrugated Fill",
    "hyd": "Pumping & Fan Hydraulics",
    "mic": "Disinfection & Microbial Risk",
    "fin": "Financial & LCC Analysis",
    "src": "Source Library",
    "mk": "Merkel Solver",
}
SCEN = ("b", "o")
MASS_SPECIES = tuple(s for s in SPECIES if s != "ph")
QMRA_INGESTION = ("rotavirus", "giardia", "cryptosporidium", "vibrio_cholerae")
_PLACEHOLDER = re.compile(r"\{([A-Za-z0-9_]+)\}")

Content = Union[str, float, int, None]


class Book:
    """Workbook + registry of named cell references + deferred formulas."""

    def __init__(self) -> None:
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.refs: Dict[str, str] = {}
        self.pending: List[Tuple[Worksheet, str, str]] = []
        self.checks: Dict[str, List[str]] = {}

    def sheet(self, key: str) -> "Sheet":
        return Sheet(self, self.wb.create_sheet(SHEETS[key]))

    def name(self, name: str, ws: Worksheet, coord: str) -> None:
        if name in self.refs:
            raise KeyError(f"duplicate reference name: {name}")
        col, row = re.match(r"([A-Z]+)(\d+)", coord).groups()   # type: ignore[union-attr]
        self.refs[name] = f"'{ws.title}'!${col}${row}"

    def name_range(self, name: str, ws: Worksheet, first: str, last: str) -> None:
        c1, r1 = re.match(r"([A-Z]+)(\d+)", first).groups()     # type: ignore[union-attr]
        c2, r2 = re.match(r"([A-Z]+)(\d+)", last).groups()      # type: ignore[union-attr]
        self.refs[name] = f"'{ws.title}'!${c1}${r1}:${c2}${r2}"

    def resolve(self, template: str) -> str:
        def sub(m: "re.Match[str]") -> str:
            key = m.group(1)
            if key not in self.refs:
                raise KeyError(f"unresolved reference {{{key}}} in {template}")
            return self.refs[key]
        return _PLACEHOLDER.sub(sub, template)

    def finalize(self) -> None:
        for ws, coord, tmpl in self.pending:
            ws[coord] = expand_sci_literals(self.resolve(tmpl))
        for title, cells in self.checks.items():
            ws = self.wb[title]
            for c in cells:
                ws.conditional_formatting.add(c, FormulaRule(formula=[f'ISNUMBER(SEARCH("❌",{c}))'], fill=CHECK_BAD_FILL))
                ws.conditional_formatting.add(c, FormulaRule(formula=[f'ISNUMBER(SEARCH("✅",{c}))'], fill=CHECK_OK_FILL))


class Sheet:
    """Row-oriented writer: B = label, C.. = values, then unit and note."""

    def __init__(self, book: Book, ws: Worksheet) -> None:
        self.book, self.ws = book, ws
        self.r = 1
        ws.sheet_view.showGridLines = True

    # -- primitives ----------------------------------------------------------------
    def put(self, coord: str, content: Content, kind: str = "out", fmt: Optional[str] = None,
            name: Optional[str] = None, bold: bool = False) -> str:
        c = self.ws[coord]
        if isinstance(content, str) and content.startswith("="):
            self.book.pending.append((self.ws, coord, content))
            c.value = "=0"
        else:
            c.value = content
        if kind == "in":
            c.fill, c.font, c.border = INPUT_FILL, INPUT_FONT, BORDER
        elif kind == "out":
            c.fill, c.font, c.border = OUTPUT_FILL, Font(bold=bold), BORDER
        elif kind == "head":
            c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        elif kind == "label":
            c.font, c.border = Font(bold=bold), BORDER
        elif kind == "note":
            c.font = NOTE_FONT
        elif kind == "unit":
            c.font, c.border = Font(color="595959"), BORDER
        if fmt:
            c.number_format = fmt
        if name:
            self.book.name(name, self.ws, coord)
        if kind == "check":
            c.border, c.font = BORDER, Font(bold=True)
            self.book.checks.setdefault(self.ws.title, []).append(coord)
        return coord

    def title(self, text: str, subtitle: str) -> None:
        self.put("B1", text, "plain")
        self.ws["B1"].font = TITLE_FONT
        self.put("B2", subtitle, "plain")
        self.ws["B2"].font = SUB_FONT
        self.r = 4

    def section(self, text: str, span: int = 6) -> None:
        self.r += 1
        for i in range(span):
            cell = self.ws.cell(row=self.r, column=2 + i)
            cell.fill = SECTION_FILL
        self.ws.cell(row=self.r, column=2, value=text).font = SECTION_FONT
        self.r += 1

    def header(self, labels: Sequence[str], col: int = 2) -> None:
        for i, lab in enumerate(labels):
            self.put(f"{get_column_letter(col + i)}{self.r}", lab, "head")
        self.r += 1

    # -- rows ----------------------------------------------------------------------
    def line(self, label: str, name: Optional[str], content: Content, unit: str = "",
             fmt: Optional[str] = None, note: str = "", kind: Optional[str] = None) -> str:
        """Single-value row: B label | C value | D unit | E note."""
        k = kind or ("out" if isinstance(content, str) and str(content).startswith("=") else "in")
        self.put(f"B{self.r}", label, "label")
        coord = self.put(f"C{self.r}", content, k, fmt, name)
        self.put(f"D{self.r}", unit, "unit")
        if note:
            self.put(f"E{self.r}", note, "note")
        self.r += 1
        return coord

    def pair(self, label: str, name: Optional[str], content: Union[Content, Tuple[Content, Content]],
             unit: str = "", fmt: Optional[str] = None, note: str = "", kind: Optional[str] = None) -> None:
        """Two-scenario row: B label | C Baseline | D Optimized | E unit | F note."""
        self.put(f"B{self.r}", label, "label")
        for i, s in enumerate(SCEN):
            if isinstance(content, tuple):
                val = content[i]
            elif isinstance(content, str) and content.startswith("="):
                val = content.replace("$", "_" + s)
            else:
                val = content
            k = kind or ("out" if isinstance(val, str) and str(val).startswith("=") else "in")
            self.put(f"{'CD'[i]}{self.r}", val, k, fmt, f"{name}_{s}" if name else None)
        self.put(f"E{self.r}", unit, "unit")
        if note:
            self.put(f"F{self.r}", note, "note")
        self.r += 1

    def widths(self, widths: Mapping[str, float]) -> None:
        for col, w in widths.items():
            self.ws.column_dimensions[col].width = w


_SCI = re.compile(r"(?<![A-Za-z0-9_$.])(\d+(?:\.\d+)?)E([+-]?\d+)")


def expand_sci_literals(formula: str) -> str:
    """Write numeric literals such as 1E-9 as plain decimals (outside string literals)."""
    parts = formula.split('"')
    for i in range(0, len(parts), 2):
        parts[i] = _SCI.sub(lambda m: format(float(m.group(0)), "f").rstrip("0").rstrip(".")
                            if abs(float(m.group(0))) >= 1 else format(float(m.group(0)), ".15f").rstrip("0"), parts[i])
    return '"'.join(parts)


# --- Formula fragments ------------------------------------------------------------------
def x_log(x: str) -> str:
    return f"LOG(MAX({x},0.001))"


def x_phs(tds: str, temp_c: str, ca: str, alk: str) -> str:
    """pHs = (9.3 + A + B) − (C + D) as an Excel expression."""
    return (f"((9.3+({x_log(tds)}-1)/10+(-13.12*LOG({temp_c}+273)+34.55))"
            f"-(({x_log(f'{ca}*2.497')}-0.4)+{x_log(alk)}))")


def x_hsat(t: str, p: str) -> str:
    """Saturated-air enthalpy h_sat(T °F) in BTU/lb dry air."""
    pw = f"(0.61094*EXP(17.625*(({t})-32)/1.8/((({t})-32)/1.8+243.04)))"
    return f"(0.24*({t})+0.62198*{pw}/({p}-{pw})*(1061+0.444*({t})))"


def x_interp(x: str, xs: str, ys: str, n: int) -> str:
    """Clamped linear interpolation in a vertical lookup table."""
    xc = f"MAX(MIN({x},INDEX({xs},{n})),INDEX({xs},1))"
    i = f"MATCH({xc},{xs},1)"
    j = f"MIN({i}+1,{n})"
    return (f"(INDEX({ys},{i})+(INDEX({ys},{j})-INDEX({ys},{i}))*({xc}-INDEX({xs},{i}))"
            f"/MAX(INDEX({xs},{j})-INDEX({xs},{i}),1E-9))")


def sel(name: str) -> str:
    """Pick the active scenario's value (Dashboard selector)."""
    return f"IF({{scen_sel}}={{lab_b}},{{{name}_b}},{{{name}_o}})"


# =========================================================================================
# Sheet builders
# =========================================================================================
def build_source_library(book: Book, cfg: Mapping[str, Any]) -> None:
    sh = book.sheet("src")
    sh.title("Source Library — Raw Water Quality Matrices",
             "Edit any yellow cell; the Dashboard source selector pulls the matching column.")
    sources = cfg["library"]["sources"]
    keys = list(sources)
    sh.section("Baseline quality by source", span=4 + len(keys))
    hdr = sh.r
    sh.header(["Parameter", "Unit"] + [sources[k]["label"] for k in keys] + ["Selected source"])
    first_col, last_col = "D", get_column_letter(3 + len(keys))
    sel_col = get_column_letter(4 + len(keys))
    book.name_range("src_labels", sh.ws, f"D{hdr}", f"{last_col}{hdr}")
    rows: List[Tuple[str, str, str, str]] = (
        [(s, SPECIES_LABELS[s], SPECIES_UNITS[s], "chemistry") for s in SPECIES]
        + [(p, PATHOGEN_LABELS[p], PATHOGEN_UNITS[p], "pathogens") for p in PATHOGEN_LABELS]
        + [("purchased", "Purchased (utility tariff) 1/0", "-", "flag"),
           ("recycled_nonpotable", "Recycled non-potable (LEED) 1/0", "-", "flag")]
    )
    for key, label, unit, group in rows:
        r = sh.r
        sh.put(f"B{r}", label, "label")
        sh.put(f"C{r}", unit, "unit")
        for j, sk in enumerate(keys):
            src = sources[sk]
            v = src[group][key] if group != "flag" else int(bool(src.get(key)))
            sh.put(f"{get_column_letter(4 + j)}{r}", float(v) if group != "flag" else v, "in",
                   F_SCI if group == "pathogens" else F_2)
        name = {"chemistry": f"raw_{key}", "pathogens": f"rawp_{key}"}.get(group, f"src_{key}")
        sh.put(f"{sel_col}{r}", f"=INDEX({first_col}{r}:{last_col}{r},MATCH({{src_sel}},{{src_labels}},0))",
               "out", F_SCI if group == "pathogens" else F_2, name)
        sh.r += 1
    sh.widths({"A": 2, "B": 32, "C": 12, sel_col: 16, **{get_column_letter(4 + j): 15 for j in range(len(keys))}})
    sh.ws.freeze_panes = f"D{hdr + 1}"


def build_dashboard(book: Book, cfg: Mapping[str, Any]) -> Sheet:
    sh = book.sheet("dash")
    sc = cfg["scenarios"]
    sh.title("Data Center Cooling Water Management — M&E Dashboard",
             f"{cfg['data_center']['it_load_mw']:.0f} MW IT design case · US customary units · "
             "yellow = input, green = live formula")
    sh.section("Scenario & source controls")
    src_label = cfg["library"]["sources"][cfg["site"]["source"]]["label"]
    sh.line("Water source", "src_sel", src_label, "", note="Dropdown — pulls Source Library column")
    sh.line("Active scenario (Mass Balance view)", "scen_sel", sc["optimized"]["label"], "",
            note="Dropdown — drives the Process Flow Mass Balance sheet")
    sh.line("IT load", "it_mw", float(cfg["data_center"]["it_load_mw"]), "MW", F_1)
    sh.line("Design wet-bulb temperature", "wb_f", float(cfg["site"]["wet_bulb_f"]), "°F", F_1)
    sh.line("Average IT utilization", "util", float(cfg["site"]["utilization"]), "-", F_PCT)
    sh.line("RO / IX demineralization on (1/0)", "ro_on", int(cfg["treatment"]["ro"]["enabled"]), "-",
            note="Skippable demineralization step")

    sh.section("Scenario levers")
    sh.put(f"B{sh.r}", "Lever", "head")
    sh.put(f"C{sh.r}", sc["baseline"]["label"], "head", name="lab_b")
    sh.put(f"D{sh.r}", sc["optimized"]["label"], "head", name="lab_o")
    sh.put(f"E{sh.r}", "Unit", "head")
    sh.r += 1
    sh.pair("Target cycles of concentration (2–10)", "coc_target",
            (float(sc["baseline"]["coc_target"]), float(sc["optimized"]["coc_target"])), "cycles", F_1,
            "Clipped by LEED/LSI/drift limits")
    sh.pair("Blowdown recovery RO on (1/0)", "bdrec",
            (int(sc["baseline"]["blowdown_recovery"]), int(sc["optimized"]["blowdown_recovery"])), "-")
    return sh


def _cell(book: Book, name: str) -> str:
    """'Sheet'!$C$6 → C6"""
    return book.refs[name].split("!")[1].replace("$", "")


def add_dashboard_validation(book: Book, sh: Sheet, cfg: Mapping[str, Any]) -> None:
    sc = cfg["scenarios"]
    dv_src = DataValidation(type="list", formula1=book.refs["src_labels"], allow_blank=False)
    dv_scn = DataValidation(type="list", formula1=f'"{sc["baseline"]["label"]},{sc["optimized"]["label"]}"')
    dv_bin = DataValidation(type="list", formula1='"1,0"')
    for dv in (dv_src, dv_scn, dv_bin):
        sh.ws.add_data_validation(dv)
    dv_src.add(_cell(book, "src_sel"))
    dv_scn.add(_cell(book, "scen_sel"))
    for nm in ("ro_on", "bdrec_b", "bdrec_o"):
        dv_bin.add(_cell(book, nm))


def finish_dashboard(sh: Sheet) -> None:
    """KPIs, checks, financial summary and charts (needs all names to exist)."""
    sh.section("Key performance indicators")
    sh.header(["KPI", "Current Baseline", "Optimized", "Δ (Opt − Base)", "Unit"])
    kpis = [
        ("Cooling load", "qcool_mw", "MW", F_2), ("Cooling load", "qcool_tons", "tons", F_INT),
        ("Tower heat-rejection load", "tl_tons", "tons", F_INT), ("Chiller COP", "cop", "-", F_2),
        ("Tower approach", "approach", "°F", F_2), ("Cycles of concentration (operating)", "coc", "cycles", F_2),
        ("Raw water intake", "intake", "gpd", F_INT), ("Raw water intake", "intake_m3d", "m³/d", F_INT),
        ("Evaporation loss", "evap_gpd", "gpd", F_INT), ("Drift loss", "drift_gpd", "gpd", F_INT),
        ("Blowdown (bleed-off)", "bd", "gpd", F_INT), ("Total discharge to sewer", "sewer_gpd", "gpd", F_INT),
        ("WUE (water usage effectiveness)", "wue", "L/kWh IT", F_3),
        ("Cooling + water plant power", "kw_total", "kW", F_INT),
        ("Cooling share of PUE", "pue_share", "-", F_3), ("PUE (incl. other overhead)", "pue", "-", F_3),
        ("Water-system energy intensity", "w_kwh_m3", "kWh/m³", F_3),
        ("Annual operating cost (incl. carbon)", "cost_total", "$/yr", F_USD),
        ("CO₂e emissions", "co2", "t/yr", F_INT), ("LEED v4 WE cooling-tower points", "leed_pts", "pts", F_INT),
        ("Legionella annual inhalation risk", "leg_py", "per yr", F_SCI),
    ]
    for label, nm, unit, fmt in kpis:
        r = sh.r
        sh.put(f"B{r}", label, "label")
        sh.put(f"C{r}", f"={{{nm}_b}}", "out", fmt)
        sh.put(f"D{r}", f"={{{nm}_o}}", "out", fmt)
        sh.put(f"E{r}", f"=D{r}-C{r}", "out", fmt)
        sh.put(f"F{r}", unit, "unit")
        sh.r += 1

    sh.section("Compliance & engineering checks")
    sh.header(["Check", "Current Baseline", "Optimized", "", "Basis"])
    checks = [
        ("CoC limiting constraint", "=\"Limited by \"&{lim$}", "LEED / LSI / drift / target"),
        ("CoC feasible with this source", "=IF({coc_raw$}>={coc_min},\"✅ Feasible\",\"❌ Enable RO/IX\")", "≥ min CoC"),
        ("Recirc silica", "=IF({cr_sio2$}>{silica_alert},\"❌ Blowdown alert\",\"✅ \"&TEXT({cr_sio2$},\"0\")&\" mg/L\")", ">150 mg/L alert"),
        ("Recirc LSI (hot water)", "=IF({lsi_hot$}>{max_lsi},\"❌ \",\"✅ \")&TEXT({lsi_hot$},\"0.00\")", "≤ inhibitor ceiling"),
        ("LEED condenser limits", "=IF({leed_ok$}=1,\"✅ All within\",\"❌ Exceeded\")", "LEED v4 Table 1"),
        ("Legionella Py ≤ target", "=IF({leg_py$}<={target_py},\"✅ \",\"❌ \")&TEXT({leg_py$},\"0.0E+00\")", "1e-4 /yr"),
        ("Ingestion QMRA (all pathogens)", "=IF({qmra_ok$}=1,\"✅ Py ≤ target\",\"❌ LRV < LRT\")", "Beta-Poisson / exp."),
        ("Pumps in HI 9.6.3 POR", "=IF({por_ok$}=1,\"✅ In POR\",\"❌ Outside POR\")", "0.7–1.2 × Q_BEP"),
        ("System water balance closes", "=IF(ABS({intake$}-{evap_gpd$}-{drift_gpd$}-{sewer_gpd$})<1,\"✅ Closed\",\"❌ Error\")", "Intake = E + D + sewer"),
    ]
    for label, tmpl, basis in checks:
        r = sh.r
        sh.put(f"B{r}", label, "label")
        sh.put(f"C{r}", tmpl.replace("$", "_b"), "check")
        sh.put(f"D{r}", tmpl.replace("$", "_o"), "check")
        sh.put(f"F{r}", basis, "note")
        sh.r += 1

    sh.section("Upgrade investment (Baseline → Optimized)")
    sh.line("CapEx (fill + eliminators + BD-RO + install)", None, "={capex_total}", "$", F_USD)
    sh.line("Year-1 net operating savings", None, "={savings}", "$/yr", F_USD)
    sh.line("NPV", None, "={npv}", "$", F_USD)
    sh.line("IRR", None, "={irr}", "-", F_PCT)
    sh.line("Discounted break-even (BEP)", None, "={bep}", "years", F_2)
    sh.line("Simple payback (SPBT)", None, "={spbt}", "years", F_2)

    sh.section("Water losses breakdown")
    sh.header(["Stream", "Current Baseline", "Optimized", "", "Unit"])
    wl_first = sh.r
    for label, nm in (("Evaporation", "evap_gpd"), ("Drift", "drift_gpd"), ("Blowdown / brine to sewer", "brine"),
                      ("Treatment rejects", "treat_waste")):
        r = sh.r
        sh.put(f"B{r}", label, "label")
        sh.put(f"C{r}", f"={{{nm}_b}}", "out", F_INT)
        sh.put(f"D{r}", f"={{{nm}_o}}", "out", F_INT)
        sh.put(f"F{r}", "gpd", "unit")
        sh.r += 1
    chart = BarChart()
    chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
    chart.title, chart.y_axis.title = "Water losses (gpd)", "gpd"
    chart.add_data(Reference(sh.ws, min_col=3, max_col=4, min_row=wl_first - 1, max_row=wl_first + 3), titles_from_data=True)
    chart.set_categories(Reference(sh.ws, min_col=2, min_row=wl_first, max_row=wl_first + 3))
    chart.height, chart.width = 8, 14
    sh.ws.add_chart(chart, "H5")

    sh.section("Annual cost breakdown")
    sh.header(["Cost item", "Current Baseline", "Optimized", "", "Unit"])
    cb_first = sh.r
    for label, nm in (("Water purchase", "cost_water"), ("Sewer", "cost_sewer"), ("Electricity", "cost_elec"),
                      ("Chemicals", "cost_chem"), ("Sludge disposal", "cost_sludge"), ("Carbon (SCC)", "cost_carbon")):
        r = sh.r
        sh.put(f"B{r}", label, "label")
        sh.put(f"C{r}", f"={{{nm}_b}}", "out", F_USD)
        sh.put(f"D{r}", f"={{{nm}_o}}", "out", F_USD)
        sh.put(f"F{r}", "$/yr", "unit")
        sh.r += 1
    bar = BarChart()
    bar.type, bar.grouping, bar.overlap = "bar", "clustered", 0
    bar.title = "Annual cost by item ($/yr)"
    bar.add_data(Reference(sh.ws, min_col=3, max_col=4, min_row=cb_first - 1, max_row=cb_first + 5), titles_from_data=True)
    bar.set_categories(Reference(sh.ws, min_col=2, min_row=cb_first, max_row=cb_first + 5))
    bar.height, bar.width = 8, 14
    sh.ws.add_chart(bar, "H22")
    sh.widths({"A": 2, "B": 42, "C": 24, "D": 26, "E": 16, "F": 22})


def build_mass_balance(book: Book, cfg: Mapping[str, Any]) -> None:
    sh = book.sheet("mb")
    tr = cfg["treatment"]
    sh.title("Process Flow Mass Balance — node by node",
             "Q_out = Q_in − Q_reject − Q_evap ;  C_out,j = [Q_in·C_in,j·(1 − R_j) + Dosing_j] / Q_out")
    sh.put("D2", "=\"Active scenario: \"&{scen_sel}", "plain")
    sh.ws["D2"].font = Font(bold=True, color="C00000")

    # Node configuration --------------------------------------------------------------
    sh.section("1. Node configuration (toggle steps on/off)", span=12)
    sh.header(["Process step", "On (1/0)", "Recovery Q_out/Q_in", "Energy kWh/kgal", "LRV virus",
               "LRV protozoa", "LRV bacteria", "pH out (blank = pass)", "Effective recovery",
               "Relative inflow Q_in/Q_intake", "Energy used kWh/kgal"])
    node_rows: Dict[str, int] = {}
    for i, key in enumerate(TREATMENT_ORDER):
        n, r = tr[key], sh.r
        node_rows[key] = r
        sh.put(f"B{r}", n["label"], "label", bold=True)
        sh.put(f"C{r}", "={ro_on}" if key == "ro" else int(n["enabled"]), "out" if key == "ro" else "in",
               None, f"on_{key}")
        sh.put(f"D{r}", float(n["recovery"]), "in", F_3, f"rec_{key}")
        e = n.get("energy_kwh_per_kgal")
        sh.put(f"E{r}", float(e) if e is not None else None, "in", F_2, f"e_{key}")
        for col, cls in zip("FGH", ("virus", "protozoa", "bacteria")):
            sh.put(f"{col}{r}", float((n.get("lrv") or {}).get(cls, 0.0)), "in", F_2, f"lrv_{cls}_{key}")
        sh.put(f"I{r}", n.get("ph_out"), "in", F_2, f"phout_{key}")
        sh.put(f"J{r}", f"=IF(C{r}=1,D{r},1)", "out", F_4, f"receff_{key}")
        sh.put(f"K{r}", 1.0 if i == 0 else f"=K{r - 1}*J{r - 1}", "out", F_4, f"rel_{key}")
        if key == "ro":
            sh.put(f"L{r}", f"=IF(C{r}=1,IF(ISNUMBER(E{r}),E{r},(0.3+0.1*{{cin_ro_tds}}/1000)*1000*0.00378541),0)",
                   "out", F_3, f"eeff_{key}")
        else:
            sh.put(f"L{r}", f"=IF(C{r}=1,E{r},0)", "out", F_3, f"eeff_{key}")
        sh.r += 1
    first, last = node_rows["intake"], node_rows["chlorination"]
    for nm, col in (("on_range", "C"), ("rel_range", "K"), ("eeff_range", "L"),
                    ("lrv_virus_range", "F"), ("lrv_protozoa_range", "G"), ("lrv_bacteria_range", "H")):
        book.name_range(nm, sh.ws, f"{col}{first}", f"{col}{last}")
    sh.r += 1
    sh.line("Overall train recovery Π(effective recovery)", "overall_rec", f"=PRODUCT(J{first}:J{last})", "-", F_4)
    sh.line("Treatment energy per gallon of intake", "treat_kwh_per_gal",
            "=SUMPRODUCT({eeff_range},{rel_range})/1000", "kWh/gal", F_SCI,
            "RO blank energy → 0.3 + 0.1·TDS(g/L) kWh/m³ (screening)")

    # Stokes / sedimentation -------------------------------------------------------------
    st = tr["coag_sed"]["stokes"]
    sh.section("2. Coagulation / sedimentation — Stokes' law settling", span=6)
    sh.line("Floc particle diameter", "st_d", float(st["particle_diameter_um"]), "µm", F_1)
    sh.line("Particle density ρp", "st_rhop", float(st["particle_density_kg_m3"]), "kg/m³", F_1)
    sh.line("Water density ρw", "st_rhow", float(st["water_density_kg_m3"]), "kg/m³", F_1)
    sh.line("Dynamic viscosity μ (15 °C)", "st_mu", float(st["viscosity_pa_s"]), "Pa·s", "0.00000")
    sh.line("Clarifier overflow rate", "st_sor", float(st["overflow_rate_m_h"]), "m/h", F_2)
    sh.line("Design TSS removal", "st_rdes", float(tr["coag_sed"]["removal"]["tss"]), "-", F_PCT)
    sh.line("Settling velocity v_s = g(ρp − ρw)d²/(18μ)", "st_vs",
            "=9.81*({st_rhop}-{st_rhow})*({st_d}*1E-6)^2/(18*{st_mu})", "m/s", F_SCI)
    sh.line("Settling velocity", "st_vs_mh", "={st_vs}*3600", "m/h", F_3)
    sh.line("Achieved TSS removal = R_design × MIN(1, v_s/SOR)", "coag_rtss",
            "={st_rdes}*MIN(1,{st_vs_mh}/{st_sor})", "-", F_PCT)
    sh.line("Alum dose", "alum_dose", float(tr["coag_sed"]["alum_dose_mg_l"]), "mg/L", F_1)

    # Removal & dosing tables ------------------------------------------------------------
    sh.section("3. Removal fraction R_j (mass to waste)  — RO row = 1 − Y·(1 − rejection)", span=16)
    sh.header(["Process step"] + [SPECIES_LABELS[s] for s in MASS_SPECIES])
    ro_rej = tr["ro"]["rejection"]
    r_rej = sh.r
    sh.put(f"B{r_rej}", "RO membrane rejection (input)", "label")
    for j, s in enumerate(MASS_SPECIES):
        sh.put(f"{get_column_letter(3 + j)}{r_rej}", float(ro_rej.get(s, 0.0)), "in", F_4, f"rej_ro_{s}")
    sh.r += 1
    for key in TREATMENT_ORDER:
        r = sh.r
        sh.put(f"B{r}", tr[key]["label"], "label")
        for j, s in enumerate(MASS_SPECIES):
            coord = f"{get_column_letter(3 + j)}{r}"
            if key == "ro":
                sh.put(coord, f"=1-{{rec_ro}}*(1-{{rej_ro_{s}}})", "out", F_4, f"R_{key}_{s}")
            elif key == "coag_sed" and s == "tss":
                sh.put(coord, "={coag_rtss}", "out", F_4, f"R_{key}_{s}")
            else:
                sh.put(coord, float((tr[key].get("removal") or {}).get(s, 0.0)), "in", F_4, f"R_{key}_{s}")
        sh.r += 1

    sh.section("4. Chemical dosing (mg/L of node inflow; negative = consumed)", span=16)
    sh.header(["Process step"] + [SPECIES_LABELS[s] for s in MASS_SPECIES])
    for key in TREATMENT_ORDER:
        r = sh.r
        sh.put(f"B{r}", tr[key]["label"], "label")
        for j, s in enumerate(MASS_SPECIES):
            sh.put(f"{get_column_letter(3 + j)}{r}", float((tr[key].get("dosing") or {}).get(s, 0.0)),
                   "in", F_2, f"D_{key}_{s}")
        sh.r += 1

    # Main matrix -------------------------------------------------------------------------
    sh.section("5. Flow & quality matrix (active scenario) — C_out,j = [Q_in·C_in,j·(1−R_j) + Dosing_j]/Q_out "
               "with Q per unit intake", span=38)
    hdr = (["Process step", "On", "Inflow Q_in (gpd)"] + [f"In {SPECIES_LABELS[s]}" for s in SPECIES]
           + ["Reject / losses (gpd)", "Evaporation (gpd)", "Outflow Q_out (gpd)"]
           + [f"Out {SPECIES_LABELS[s]}" for s in SPECIES] + ["Flow balance check", "TSS removed (kg/gal intake)"])
    sh.header(hdr)
    sh.ws.row_dimensions[sh.r - 1].height = 45
    cin_col = {s: get_column_letter(5 + i) for i, s in enumerate(SPECIES)}          # E..S
    cout_col = {s: get_column_letter(23 + i) for i, s in enumerate(SPECIES)}        # W..AK
    mat_first = sh.r
    prev: Optional[str] = None
    for key in TREATMENT_ORDER:
        r = sh.r
        sh.put(f"B{r}", tr[key]["label"], "label", bold=True)
        sh.put(f"C{r}", f"={{on_{key}}}", "out")
        q_in = f"={sel('mf')}/{{overall_rec}}" if prev is None else f"={{qout_{prev}}}"
        sh.put(f"D{r}", q_in, "out", F_INT, f"qin_{key}")
        for s in SPECIES:
            src = f"={{raw_{s}}}" if prev is None else f"={{cout_{prev}_{s}}}"
            sh.put(f"{cin_col[s]}{r}", src, "out", F_2, f"cin_{key}_{s}")
        sh.put(f"T{r}", f"=D{r}-V{r}", "out", F_INT, f"qrej_{key}")
        sh.put(f"U{r}", 0, "out", F_INT)
        sh.put(f"V{r}", f"=D{r}*{{receff_{key}}}", "out", F_INT, f"qout_{key}")
        for s in SPECIES:
            cin = f"{cin_col[s]}{r}"
            if s == "ph":
                f = f"=IF(AND(C{r}=1,ISNUMBER({{phout_{key}}})),{{phout_{key}}},{cin})"
            else:
                # Same equation with flows normalised per unit intake (Q_in,rel = K, Q_out,rel = K·J):
                # concentrations stay flow-independent, so Excel has no circular reference.
                q_in, q_out = f"{{rel_{key}}}", f"({{rel_{key}}}*{{receff_{key}}})"
                f = (f"=IF(C{r}=1,MAX(0,({q_in}*{cin}*(1-{{R_{key}_{s}}})+{{D_{key}_{s}}}*{q_in})/{q_out}),{cin})")
            sh.put(f"{cout_col[s]}{r}", f, "out", F_2, f"cout_{key}_{s}")
        sh.put(f"AL{r}", f"=IF(ABS(D{r}-T{r}-U{r}-V{r})<0.001,\"✅\",\"❌\")", "check")
        sh.put(f"AM{r}", f"=IF(C{r}=1,{{rel_{key}}}*{cin_col['tss']}{r}*{{R_{key}_tss}},0)*3.78541/1E6", "out", F_SCI)
        prev = key
        sh.r += 1
    book.name_range("tss_removed_range", sh.ws, f"AM{mat_first}", f"AM{sh.r - 1}")
    # Cooling tower row
    r = sh.r
    sh.put(f"B{r}", "Cooling Tower (recirculation → blowdown)", "label", bold=True)
    sh.put(f"C{r}", 1, "out")
    sh.put(f"D{r}", f"={sel('mt')}", "out", F_INT)
    for s in SPECIES:
        sh.put(f"{cin_col[s]}{r}", f"={sel('cbl_' + s)}", "out", F_2)
    sh.put(f"T{r}", f"={sel('drift_gpd')}", "out", F_INT)
    sh.put(f"U{r}", f"={sel('evap_gpd')}", "out", F_INT)
    sh.put(f"V{r}", f"={sel('bd')}", "out", F_INT)
    for s in SPECIES:
        sh.put(f"{cout_col[s]}{r}", f"={sel('cr_' + s)}", "out", F_2)
    sh.put(f"AL{r}", f"=IF(ABS(D{r}-T{r}-U{r}-V{r})<0.01,\"✅\",\"❌\")", "check")
    sh.r += 1
    # Blowdown recovery row
    r = sh.r
    sh.put(f"B{r}", "Blowdown Recovery RO (permeate → makeup)", "label", bold=True)
    sh.put(f"C{r}", f"={sel('bdrec')}", "out")
    sh.put(f"D{r}", f"=IF(C{r}=1,{sel('bd')},0)", "out", F_INT)
    for s in SPECIES:
        sh.put(f"{cin_col[s]}{r}", f"={sel('cr_' + s)}", "out", F_2)
    sh.put(f"T{r}", f"=D{r}-V{r}", "out", F_INT)
    sh.put(f"U{r}", 0, "out", F_INT)
    sh.put(f"V{r}", f"=IF(C{r}=1,{sel('perm')},0)", "out", F_INT)
    for s in SPECIES:
        sh.put(f"{cout_col[s]}{r}", f"=IF(C{r}=1,{sel('cp_' + s)},0)", "out", F_2)
    sh.put(f"AL{r}", f"=IF(ABS(D{r}-T{r}-U{r}-V{r})<0.01,\"✅\",\"❌\")", "check")
    sh.r += 2
    sh.line("System check: Intake − Evaporation − Drift − Sewer discharges", None,
            f"={sel('intake')}-{sel('evap_gpd')}-{sel('drift_gpd')}-{sel('sewer_gpd')}", "gpd", F_2,
            "Must be 0 — whole-plant water balance")
    sh.widths({"A": 2, "B": 44, "C": 10, "D": 16, **{c: 11 for c in list(cin_col.values()) + list(cout_col.values())},
               "T": 16, "U": 15, "V": 16, "AL": 12, "AM": 14, "E": 12, "F": 12, "G": 12, "H": 12, "I": 14,
               "J": 14, "K": 16, "L": 14})


def build_cooling(book: Book, cfg: Mapping[str, Any]) -> None:
    sh = book.sheet("ct")
    sc, ct, dc, fm = cfg["scenarios"], cfg["cooling_tower"], cfg["data_center"], cfg["fill_model"]
    sh.title("Cooling Tower Heat Balance & Corrugated Fill Simulation",
             "Q = ṁ·cp·ΔT · Tower Load = Q(1 + 1/COP) · E = TL/(970×8.34) · B = (E − (N−1)D)/(N·k − 1 + Y)")

    sh.section("1. Shared design inputs")
    sh.line("Other heat to chilled water (fraction of IT)", "other_heat", float(dc["other_heat_fraction"]), "-", F_PCT)
    sh.line("Chilled-water ΔT", "chw_dt", float(dc["chw_delta_t_f"]), "°F", F_1)
    sh.line("Chiller COP at reference CWST", "cop_ref", float(dc["cop_ref"]), "-", F_2)
    sh.line("Reference condenser-water supply temp", "cwst_ref", float(dc["cwst_ref_f"]), "°F", F_1)
    sh.line("COP sensitivity per °F colder CWST", "cop_sens", float(dc["cop_sensitivity_per_f"]), "1/°F", F_PCT,
            "Screening ~2 %/°F")
    sh.line("Non-cooling PUE overhead", "pue_other", float(dc["other_overhead_pue"]), "-", F_3)
    sh.line("Tower range (hot − cold)", "range_f", float(ct["range_f"]), "°F", F_1)
    sh.line("Water/air mass ratio L/G", "lg", float(ct["l_over_g"]), "-", F_2)
    sh.line("Barometric pressure", "p_kpa", float(cfg["site"]["barometric_kpa"]), "kPa", F_3)
    sh.line("Minimum CoC", "coc_min", float(ct["coc_min"]), "cycles", F_2)
    sh.line("Max recirculating LSI (inhibitor program)", "max_lsi", float(ct["max_lsi"]), "-", F_2)
    sh.line("Recirculating free-chlorine residual", "recirc_cl2", float(ct["recirc_cl2_mg_l"]), "mg/L", F_2)
    sh.line("Silica blowdown alert", "silica_alert", float(cfg["library"]["silica_alert_mg_l"]), "mg/L", F_INT)
    sh.line("Blowdown-recovery RO recovery Y", "bd_y", float(cfg["blowdown_recovery"]["recovery"]), "-", F_PCT)
    lim = cfg["library"]["leed_condenser_limits"]
    for k, lab, unit in (("ca_as_caco3", "LEED max Ca (as CaCO3)", "ppm"), ("alkalinity", "LEED max alkalinity", "ppm"),
                         ("sio2", "LEED max SiO₂", "ppm"), ("cl", "LEED max chloride", "ppm"),
                         ("ec", "LEED max conductivity", "µS/cm")):
        sh.line(lab, f"lim_{k}", float(lim[k]), unit, F_INT, "LEED v4 WE Cooling Tower Water Use, Table 1")
    sh.line("Fill model c0 (KaV/L at reference)", "fm_c0", float(fm["c0"]), "-", F_3, "Screening calibration")
    sh.line("Reference SSA", "fm_ssa", float(fm["ssa_ref_m2_m3"]), "m²/m³", F_1)
    sh.line("Reference fill height", "fm_h", float(fm["height_ref_m"]), "m", F_2)
    sh.line("KaV/L exponent on L/G (n)", "fm_n", float(fm["lg_exponent"]), "-", F_2)
    sh.line("Flute-angle thermal exponent", "fm_a", float(fm["angle_exponent"]), "-", F_2)

    sh.section("2. Scenario inputs — fill, drift eliminator, chemistry")
    sh.header(["Parameter", sc["baseline"]["label"], sc["optimized"]["label"], "Unit", "Note"])
    fb, fo = sc["baseline"]["fill"], sc["optimized"]["fill"]
    sh.pair("Fill / eliminator description", "fill_name", (fb["name"], fo["name"]), "")
    sh.pair("Specific surface area (SSA)", "ssa", (float(fb["ssa_m2_m3"]), float(fo["ssa_m2_m3"])), "m²/m³", F_1)
    sh.pair("Flute angle", "ang", (float(fb["flute_angle_deg"]), float(fo["flute_angle_deg"])), "°", F_1)
    sh.pair("Sheet thickness", "thk", (float(fb["sheet_thickness_mm"]), float(fo["sheet_thickness_mm"])), "mm", F_2)
    sh.pair("Fill height", "fh", (float(fb["fill_height_m"]), float(fo["fill_height_m"])), "m", F_2)
    sh.pair("Drift fraction of recirculation", "fdrift", (float(fb["drift_fraction"]), float(fo["drift_fraction"])),
            "-", F_PCT3, "0.20 % → 0.001 % with HE eliminator")
    sh.pair("Drift eliminator pressure drop", "de_dp", (float(fb["eliminator_dp_pa"]), float(fo["eliminator_dp_pa"])), "Pa", F_1)
    sh.pair("Recirculating pH", "ph_r", (float(sc["baseline"]["recirc_ph"]), float(sc["optimized"]["recirc_ph"])), "-", F_2)
    sh.pair("Target CoC (from Dashboard)", "coc_t", "={coc_target$}", "cycles", F_2)
    sh.pair("Blowdown recovery on (from Dashboard)", "bdr", "={bdrec$}", "-")

    sh.section("3. Heat balance")
    sh.pair("Cooling load Q = IT × 3412 × (1 + other)", "qcool", "={it_mw}*1000*3412*(1+{other_heat})", "BTU/hr", F_INT)
    sh.pair("Cooling load", "qcool_mw", "={qcool$}/3412/1000", "MW", F_2)
    sh.pair("Cooling load", "qcool_tons", "={qcool$}/12000", "tons", F_INT)
    sh.pair("Fill KaV/L available = c0(SSA/SSAref)(H/Href)(θ/30)^a(L/G)^−n", "kavl",
            "={fm_c0}*({ssa$}/{fm_ssa})*({fh$}/{fm_h})*({ang$}/30)^{fm_a}*{lg}^(-{fm_n})", "-", F_3)
    mk_idx = "MATCH({kavl$},{mk_req},-1)"
    sh.pair("Cold-water temperature (Merkel solver, interpolated)", "tcold",
            f"=IFERROR(INDEX({{mk_tc}},{mk_idx})+(INDEX({{mk_tc}},{mk_idx}+1)-INDEX({{mk_tc}},{mk_idx}))"
            f"*(INDEX({{mk_req}},{mk_idx})-{{kavl$}})/(INDEX({{mk_req}},{mk_idx})-INDEX({{mk_req}},{mk_idx}+1)),{{wb_f}}+0.1)",
            "°F", F_2, "Required KaV/L = available (Chebyshev)")
    sh.pair("Approach (cold water − wet bulb)", "approach", "={tcold$}-{wb_f}", "°F", F_2)
    sh.pair("Hot-water temperature", "thot", "={tcold$}+{range_f}", "°F", F_2)
    sh.pair("Hot-water temperature", "hot_c", "=({thot$}-32)/1.8", "°C", F_2)
    sh.pair("Cold-water temperature", "cold_c", "=({tcold$}-32)/1.8", "°C", F_2)
    sh.pair("Chiller COP = COPref·(1 + s·(CWSTref − CWST))", "cop", "={cop_ref}*(1+{cop_sens}*({cwst_ref}-{tcold$}))", "-", F_3)
    sh.pair("Chiller power", "kw_chiller", "={qcool$}/3412/{cop$}", "kW", F_INT)
    sh.pair("Tower Load = Q·(1 + 1/COP)", "tl", "={qcool$}*(1+1/{cop$})", "BTU/hr", F_INT)
    sh.pair("Tower Load", "tl_tons", "={tl$}/12000", "tons", F_INT)
    sh.pair("Recirculation Q = TL/(8.34·60·Range)", "recirc", "={tl$}/(8.34*60*{range_f})", "gpm", F_INT)
    sh.pair("Evaporation E = TL/(970 × 8.34)", "evap_gph", "={tl$}/(970*8.34)", "gph", F_INT)
    sh.pair("Evaporation", "evap_gpd", "={evap_gph$}*24", "gpd", F_INT)
    sh.pair("Drift D = f_drift × Q_recirc", "drift_gpd", "={recirc$}*{fdrift$}*1440", "gpd", F_INT)

    sh.section("4. Cycles of concentration — limits on treated makeup")
    sh.header(["Constraint", sc["baseline"]["label"], sc["optimized"]["label"], "Unit", "Basis"])
    cand_first = sh.r
    sh.pair("Target", "cand_t", "={coc_t$}", "cycles", F_2)
    sh.pair("LEED Ca (as CaCO3)", "cand_ca", "={lim_ca_as_caco3}/MAX({cf_ca}*2.497,0.001)", "cycles", F_2)
    sh.pair("LEED alkalinity", "cand_alk", "={lim_alkalinity}/MAX({cf_alkalinity},0.001)", "cycles", F_2)
    sh.pair("LEED SiO₂", "cand_si", "={lim_sio2}/MAX({cf_sio2},0.001)", "cycles", F_2)
    sh.pair("LEED chloride", "cand_cl", "={lim_cl}/MAX({cf_cl},0.001)", "cycles", F_2)
    sh.pair("LEED conductivity", "cand_ec", "={lim_ec}/MAX({cf_ec},0.001)", "cycles", F_2)
    sh.pair("LSI ceiling: 10^((LSImax − pHr + pHs₁)/1.9)", "cand_lsi",
            f"=10^(({{max_lsi}}-{{ph_r$}}+{x_phs('{cf_tds}', '{hot_c$}', '{cf_ca}', '{cf_alkalinity}')})/1.9)", "cycles", F_2)
    sh.pair("Drift (bleed ≥ 0): 1 + E/D", "cand_drift", "=IF({drift_gpd$}>0,1+{evap_gpd$}/{drift_gpd$},1E9)", "cycles", F_2)
    cand_last = sh.r - 1
    for i, s in enumerate(SCEN):
        col = "CD"[i]
        book.name_range(f"cands_{s}", sh.ws, f"{col}{cand_first}", f"{col}{cand_last}")
    book.name_range("cand_names", sh.ws, f"B{cand_first}", f"B{cand_last}")
    sh.pair("Most restrictive limit", "coc_raw", "=MIN({cands$})", "cycles", F_2)
    sh.pair("Limiting constraint", "lim", "=INDEX({cand_names},MATCH({coc_raw$},{cands$},0))", "")
    sh.pair("Operating CoC = MAX(min CoC, limit)", "coc", "=MAX({coc_min},{coc_raw$})", "cycles", F_2)

    sh.section("5. Tower & blowdown-recovery water balance")
    sh.pair("BD-RO recovery used Y", "y", "=IF({bdr$}=1,{bd_y},0)", "-", F_PCT)
    sh.pair("TDS mass fraction to brine k", "ktds_used", "={k_tds$}", "-", F_4)
    sh.pair("Blowdown B = (E − (N−1)D)/(N·k − 1 + Y)", "bd",
            "=MAX(0,({evap_gpd$}-({coc$}-1)*{drift_gpd$})/({coc$}*{k_tds$}-1+{y$}))", "gpd", F_INT)
    sh.pair("Course check: B = E/(CoC − 1)", "bd_simple", "={evap_gpd$}/({coc$}-1)", "gpd", F_INT,
            "Drift lumped into B; recovery off")
    sh.pair("Permeate recycled P = Y·B", "perm", "={y$}*{bd$}", "gpd", F_INT)
    sh.pair("Brine / blowdown to sewer", "brine", "={bd$}-{perm$}", "gpd", F_INT)
    sh.pair("Fresh makeup M_f = E + D + brine", "mf", "={evap_gpd$}+{drift_gpd$}+{brine$}", "gpd", F_INT)
    sh.pair("Tower makeup M_t = M_f + P", "mt", "={mf$}+{perm$}", "gpd", F_INT)
    sh.pair("Raw water intake = M_f / Π recovery", "intake", "={mf$}/{overall_rec}", "gpd", F_INT)
    sh.pair("Raw water intake", "intake_m3d", "={intake$}*0.00378541", "m³/d", F_INT)
    sh.pair("Treatment rejects", "treat_waste", "={intake$}-{mf$}", "gpd", F_INT)
    sh.pair("Total discharge to sewer", "sewer_gpd", "={treat_waste$}+{brine$}", "gpd", F_INT)
    sh.pair("WUE = intake / IT energy", "wue", "={intake$}*3.78541/({it_mw}*1000*24)", "L/kWh", F_3)

    sh.section("6. Recirculating chemistry  C_r,j = M_f·C_f,j / (D + B·k_j)", span=16)
    hdr = ["Species", "Treated makeup C_f", "BD-RO rejection"]
    for s in SCEN:
        tag = "Base" if s == "b" else "Opt"
        hdr += [f"k ({tag})", f"C_recirc ({tag})", f"C_permeate ({tag})", f"C_brine ({tag})", f"C_blend ({tag})"]
    sh.header(hdr)
    bd_rej = cfg["blowdown_recovery"]["rejection"]
    for s in SPECIES:
        r = sh.r
        sh.put(f"B{r}", SPECIES_LABELS[s], "label")
        sh.put(f"C{r}", f"={{cout_chlorination_{s}}}", "out", F_2, f"cf_{s}")
        if s != "ph":
            sh.put(f"D{r}", float(bd_rej.get(s, 0.0)), "in", F_4, f"bdrej_{s}")
        for i, sx in enumerate(SCEN):
            c0 = 5 + i * 5
            col = [get_column_letter(c0 + k) for k in range(5)]
            if s == "ph":
                sh.put(f"{col[1]}{r}", f"={{ph_r_{sx}}}", "out", F_2, f"cr_{s}_{sx}")
                sh.put(f"{col[2]}{r}", f"={{ph_r_{sx}}}", "out", F_2, f"cp_{s}_{sx}")
                sh.put(f"{col[3]}{r}", f"={{ph_r_{sx}}}", "out", F_2, f"cbr_{s}_{sx}")
                sh.put(f"{col[4]}{r}", f"=C{r}", "out", F_2, f"cbl_{s}_{sx}")
                continue
            sh.put(f"{col[0]}{r}", f"=IF({{bdr_{sx}}}=1,1-{{bd_y}}*(1-D{r}),1)", "out", F_4, f"k_{s}_{sx}")
            cr = ("={recirc_cl2}" if s == "cl2" else
                  f"={{mf_{sx}}}*C{r}/MAX({{drift_gpd_{sx}}}+{{bd_{sx}}}*{col[0]}{r},1E-12)")
            sh.put(f"{col[1]}{r}", cr, "out", F_2, f"cr_{s}_{sx}")
            sh.put(f"{col[2]}{r}", f"=IF({{bdr_{sx}}}=1,{col[1]}{r}*(1-{col[0]}{r})/{{bd_y}},0)", "out", F_2, f"cp_{s}_{sx}")
            sh.put(f"{col[3]}{r}", f"=IF({{bdr_{sx}}}=1,{col[1]}{r}*{col[0]}{r}/(1-{{bd_y}}),{col[1]}{r})", "out", F_2,
                   f"cbr_{s}_{sx}")
            sh.put(f"{col[4]}{r}", f"=({{mf_{sx}}}*C{r}+{{perm_{sx}}}*{col[2]}{r})/{{mt_{sx}}}", "out", F_2, f"cbl_{s}_{sx}")
        sh.r += 1

    sh.section("7. Scaling indices, silica & LEED v4 WE credit")
    sh.header(["Index", sc["baseline"]["label"], sc["optimized"]["label"], "Unit", "Status / basis"])
    sh.pair("Makeup pHs (cold water)", "phs_mu", "=" + x_phs("{cf_tds}", "{cold_c$}", "{cf_ca}", "{cf_alkalinity}"), "-", F_2)
    sh.pair("Makeup LSI = pH − pHs", "lsi_mu", "={cf_ph}-{phs_mu$}", "-", F_2)
    sh.pair("Recirc pHs (cold water)", "phs_rc", "=" + x_phs("{cr_tds$}", "{cold_c$}", "{cr_ca$}", "{cr_alkalinity$}"), "-", F_2)
    sh.pair("Recirc LSI (cold water)", "lsi_cold", "={ph_r$}-{phs_rc$}", "-", F_2)
    sh.pair("Recirc pHs (hot water)", "phs_rh", "=" + x_phs("{cr_tds$}", "{hot_c$}", "{cr_ca$}", "{cr_alkalinity$}"), "-", F_2)
    sh.pair("Recirc LSI (hot water)", "lsi_hot", "={ph_r$}-{phs_rh$}", "-", F_2)
    sh.pair("Recirc RSI = 2pHs − pH (hot water)", "rsi_hot", "=2*{phs_rh$}-{ph_r$}", "-", F_2)
    sh.pair("LSI status", "lsi_status",
            "=IF({lsi_hot$}>2,\"Severe scaling\",IF({lsi_hot$}>0.5,\"Scale-forming\",IF({lsi_hot$}>=-0.5,\"Balanced\","
            "IF({lsi_hot$}>=-2,\"Corrosive\",\"Severely corrosive\"))))", "")
    sh.pair("RSI status", "rsi_status",
            "=IF({rsi_hot$}<5.5,\"Heavy scale\",IF({rsi_hot$}<6.2,\"Scale-forming\",IF({rsi_hot$}<=6.8,\"Stable\","
            "IF({rsi_hot$}<=8.5,\"Corrosive\",\"Severely corrosive\"))))", "")
    sh.pair("Silica check", "si_chk",
            "=IF({cr_sio2$}>{silica_alert},\"❌ Blowdown alert\",\"✅ Below alert\")", "", kind="check")
    sh.pair("All LEED Table 1 limits met (1/0)", "leed_ok",
            "=IF(AND({cr_ca$}*2.497<={lim_ca_as_caco3}*1.0001,{cr_alkalinity$}<={lim_alkalinity}*1.0001,"
            "{cr_sio2$}<={lim_sio2}*1.0001,{cr_cl$}<={lim_cl}*1.0001,{cr_ec$}<={lim_ec}*1.0001),1,0)", "-")
    sh.pair("LEED max cycles (≤ 10)", "leed_max", "=MIN({cand_ca$},{cand_alk$},{cand_si$},{cand_cl$},{cand_ec$},10)", "cycles", F_2)
    sh.pair("Recycled non-potable share of makeup", "nonpot",
            "=({perm$}+IF({src_recycled_nonpotable}=1,{mf$},0))/{mt$}", "-", F_PCT)
    sh.pair("LEED v4 WE Cooling Tower Water Use points", "leed_pts",
            "=IF(OR(AND({coc$}>=10-1E-6,OR({ro_on}=1,{bdr$}=1)),AND({coc$}>={leed_max$}-1E-6,{nonpot$}>=0.2)),2,"
            "IF({coc$}>={leed_max$}-1E-6,1,0))", "pts", F_INT, "1 pt: max cycles; 2 pts: ≥10 cycles w/ treatment or ≥20 % recycled")

    sh.section("8. Corrugated fill impact — Baseline vs Upgrade")
    sh.line("Approach improvement", None, "={approach_b}-{approach_o}", "°F", F_2)
    sh.line("COP improvement", None, "={cop_o}/{cop_b}-1", "-", F_PCT)
    sh.line("Evaporation reduction", None, "={evap_gpd_b}-{evap_gpd_o}", "gpd", F_INT)
    sh.line("Drift reduction", None, "={drift_gpd_b}-{drift_gpd_o}", "gpd", F_INT)
    sh.line("Blowdown-to-sewer reduction", None, "={brine_b}-{brine_o}", "gpd", F_INT)
    sh.line("Raw water saved", None, "={intake_b}-{intake_o}", "gpd", F_INT)
    sh.line("Fan static pressure increase", None, "={dp_total_o}-{dp_total_b}", "Pa", F_1)
    sh.line("Fan power change", None, "={kw_fan_o}-{kw_fan_b}", "kW", F_1)
    sh.line("Chiller power saved", None, "={kw_chiller_b}-{kw_chiller_o}", "kW", F_1)

    chart = BarChart()
    chart.type, chart.grouping = "col", "clustered"
    chart.title, chart.y_axis.title = "Tower water losses: Baseline vs Upgrade (gpd)", "gpd"
    ws = sh.ws
    r0 = sh.r + 2
    for i, (lab, b, o) in enumerate((("Stream", "Baseline", "Optimized"), ("Evaporation", "evap_gpd_b", "evap_gpd_o"),
                                     ("Drift", "drift_gpd_b", "drift_gpd_o"), ("Sewer", "sewer_gpd_b", "sewer_gpd_o"),
                                     ("Intake", "intake_b", "intake_o"))):
        if i == 0:
            sh.put(f"H{r0}", lab, "head"); sh.put(f"I{r0}", b, "head"); sh.put(f"J{r0}", o, "head")
        else:
            sh.put(f"H{r0 + i}", lab, "label")
            sh.put(f"I{r0 + i}", f"={{{b}}}", "out", F_INT)
            sh.put(f"J{r0 + i}", f"={{{o}}}", "out", F_INT)
    chart.add_data(Reference(ws, min_col=9, max_col=10, min_row=r0, max_row=r0 + 4), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=8, min_row=r0 + 1, max_row=r0 + 4))
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, f"H{r0 + 6}")
    sh.widths({"A": 2, "B": 52, "C": 26, "D": 28, "E": 12, "F": 34, "G": 12, "H": 14, "I": 14, "J": 14,
               "K": 12, "L": 12, "M": 12, "N": 12, "O": 12, "P": 12, "Q": 12})


def build_merkel(book: Book) -> None:
    sh = book.sheet("mk")
    sh.title("Merkel Solver — required tower characteristic KaV/L vs cold-water temperature",
             "KaV/L = Range/4 · Σ 1/(h_sat(T_i) − h_air,i), T_i = T_cold + {0.1,0.4,0.6,0.9}·Range (Chebyshev)")
    sh.line("Inlet air enthalpy h_sat(T_wb)", "mk_hin", "=" + x_hsat("{wb_f}", "{p_kpa}"), "BTU/lb", F_3)
    sh.r += 1
    sh.header(["Approach (°F)", "T_cold (°F)", "Δh @0.1R", "Δh @0.4R", "Δh @0.6R", "Δh @0.9R", "Required KaV/L"])
    first = sh.r
    steps = [round(0.1 * i, 1) for i in range(1, 401)]
    for a in steps:
        r = sh.r
        sh.ws[f"B{r}"] = a
        sh.ws[f"B{r}"].number_format = F_1
        sh.put(f"C{r}", f"={{wb_f}}+B{r}", "out", F_2)
        for col, f in zip("DEFG", (0.1, 0.4, 0.6, 0.9)):
            t = f"(C{r}+{f}*{{range_f}})"
            sh.put(f"{col}{r}", f"={x_hsat(t, '{p_kpa}')}-({{mk_hin}}+{{lg}}*{f}*{{range_f}})", "out", F_3)
        sh.put(f"H{r}", f"=IF(MIN(D{r}:G{r})<=0,99999,{{range_f}}/4*(1/D{r}+1/E{r}+1/F{r}+1/G{r}))", "out", F_4)
        sh.r += 1
    book.name_range("mk_tc", sh.ws, f"C{first}", f"C{sh.r - 1}")
    book.name_range("mk_req", sh.ws, f"H{first}", f"H{sh.r - 1}")
    sh.widths({"A": 2, "B": 34, "C": 14, "D": 12, "E": 12, "F": 12, "G": 12, "H": 16})
    sh.ws.freeze_panes = f"B{first}"


def build_hydraulics(book: Book, cfg: Mapping[str, Any]) -> None:
    sh = book.sheet("hyd")
    ct, fm = cfg["cooling_tower"], cfg["fill_model"]
    sc = cfg["scenarios"]
    sh.title("Pumping & Induced-Draft Fan Hydraulics",
             "WHP = Q·TDH/3960 · BHP = WHP/η_h · kW = BHP·0.7457/(η_m·η_e) · N = ROUNDUP(Q/Q_BEP) · POR 0.7–1.2 Q_BEP")
    sh.section("1. Fan & fill air-side inputs")
    sh.line("Tower water loading", "wload", float(ct["water_loading_gpm_ft2"]), "gpm/ft²", F_1)
    sh.line("Exhaust air specific volume", "air_sv", float(ct["air_specific_volume_ft3_lb"]), "ft³/lb", F_2)
    sh.line("Air density", "air_rho", float(ct["air_density_kg_m3"]), "kg/m³", F_3)
    sh.line("Other static pressure (louvers, plenum, exit)", "dp_other", float(ct["other_static_pa"]), "Pa", F_1)
    sh.line("Fan efficiency", "eta_fan", float(ct["fan_efficiency"]), "-", F_PCT)
    sh.line("Fan motor efficiency", "eta_fmot", float(ct["fan_motor_efficiency"]), "-", F_PCT)
    sh.line("Fill ΔP coefficient c", "fm_dpc", float(fm["dp_coefficient"]), "-", F_3, "Screening calibration")
    sh.line("Fill ΔP flute-angle exponent", "fm_dpa", float(fm["dp_angle_exponent"]), "-", F_2)

    sh.section("2. Induced-draft fan energy")
    sh.header(["Parameter", sc["baseline"]["label"], sc["optimized"]["label"], "Unit", "Formula"])
    sh.pair("Tower plan area", "area", "={recirc$}/{wload}/10.7639", "m²", F_1, "Q_recirc / loading")
    sh.pair("Air mass flow G = L/(L/G)", "g_air", "={recirc$}*8.34*60/{lg}", "lb/hr", F_INT)
    sh.pair("Airflow", "cfm", "={g_air$}*{air_sv}/60", "cfm", F_INT)
    sh.pair("Airflow", "airflow", "={cfm$}*0.0283168/60", "m³/s", F_1)
    sh.pair("Fill face velocity", "vface", "={airflow$}/{area$}", "m/s", F_2)
    sh.pair("Fill open area ε = 1 − SSA·t/2", "eps", "=1-{ssa$}*{thk$}/1000/2", "-", F_4)
    sh.pair("Fill ΔP = c·SSA·H·½ρv²·(θ/30)^b/ε²", "dp_fill",
            "={fm_dpc}*{ssa$}*{fh$}*0.5*{air_rho}*{vface$}^2*({ang$}/30)^{fm_dpa}/{eps$}^2", "Pa", F_1)
    sh.pair("Total static pressure", "dp_total", "={dp_fill$}+{de_dp$}+{dp_other}", "Pa", F_1)
    sh.pair("Total static pressure", "dp_total_inwg", "={dp_total$}/249.089", "in. w.g.", F_3)
    sh.pair("IDF fan power = V·ΔP/(η_f·η_m)", "kw_fan", "={airflow$}*{dp_total$}/({eta_fan}*{eta_fmot})/1000", "kW", F_1)

    sh.section("3. Pump stations (both scenarios)", span=14)
    sh.header(["Pump station", "Scenario", "Flow Q (gpm)", "TDH (ft)", "Q_BEP (gpm)", "η hydraulic", "η mechanical",
               "η electrical", "N = ROUNDUP(Q/Q_BEP)", "Q per pump", "POR min", "POR max", "In POR?",
               "WHP", "BHP", "kW"])
    pumps = cfg["pumps"]
    flows = {"chw": "={qcool$}/(8.34*60*{chw_dt})", "cw": "={recirc$}", "makeup": "={intake$}/1440"}
    por_cells: Dict[str, List[str]] = {"b": [], "o": []}
    for key in ("chw", "cw", "makeup"):
        p = pumps[key]
        for s in SCEN:
            r = sh.r
            sh.put(f"B{r}", p["label"], "label")
            sh.put(f"C{r}", "Baseline" if s == "b" else "Optimized", "label")
            sh.put(f"D{r}", flows[key].replace("$", "_" + s), "out", F_INT)
            first_input = s == "b"
            for col, fld, fmt in (("E", "tdh_ft", F_1), ("F", "q_bep_gpm", F_INT), ("G", "eff_hydraulic", F_PCT),
                                  ("H", "eff_mechanical", F_PCT), ("I", "eff_electrical", F_PCT)):
                if first_input:
                    sh.put(f"{col}{r}", float(p[fld]), "in", fmt, f"p_{key}_{fld}")
                else:
                    sh.put(f"{col}{r}", f"={{p_{key}_{fld}}}", "out", fmt)
            sh.put(f"J{r}", f"=MAX(1,ROUNDUP(D{r}/F{r},0))", "out", F_INT)
            sh.put(f"K{r}", f"=D{r}/J{r}", "out", F_INT)
            sh.put(f"L{r}", f"=0.7*F{r}", "out", F_INT)
            sh.put(f"M{r}", f"=1.2*F{r}", "out", F_INT)
            sh.put(f"N{r}", f"=IF(AND(K{r}>=L{r},K{r}<=M{r}),\"✅ In POR\",\"❌ Outside\")", "check")
            sh.put(f"O{r}", f"=D{r}*E{r}/3960", "out", F_1)
            sh.put(f"P{r}", f"=O{r}/G{r}", "out", F_1)
            sh.put(f"Q{r}", f"=P{r}*0.7457/(H{r}*I{r})", "out", F_1, f"kw_{key}_{s}")
            por_cells[s].append(f"N{r}")
            sh.r += 1
    sh.r += 1
    sh.section("4. Plant power summary")
    sh.header(["Component", sc["baseline"]["label"], sc["optimized"]["label"], "Unit", ""])
    pw_first = sh.r
    sh.pair("Chiller compressors", "pw_chiller", "={kw_chiller$}", "kW", F_INT)
    sh.pair("Tower fans", "pw_fan", "={kw_fan$}", "kW", F_INT)
    sh.pair("Chilled-water pumps", "pw_chw", "={kw_chw$}", "kW", F_INT)
    sh.pair("Condenser-water pumps", "pw_cw", "={kw_cw$}", "kW", F_INT)
    sh.pair("Raw/makeup water pumps", "pw_mu", "={kw_makeup$}", "kW", F_1)
    sh.pair("BD-RO energy", "bdro_kwh",
            "=IF({bdr$}=1,(0.3+0.1*{cr_tds$}/1000)*1000*0.00378541*{bd$}/1000,0)", "kWh/d", F_INT)
    sh.pair("Water treatment (train + BD-RO)", "pw_treat", "=({treat_kwh_per_gal}*{intake$}+{bdro_kwh$})/24", "kW", F_1)
    sh.pair("Total cooling + water plant", "kw_total",
            "={pw_chiller$}+{pw_fan$}+{pw_chw$}+{pw_cw$}+{pw_mu$}+{pw_treat$}", "kW", F_INT)
    sh.pair("Cooling share of PUE", "pue_share", "={kw_total$}/({it_mw}*1000)", "-", F_3)
    sh.pair("PUE", "pue", "=1+{pue_share$}+{pue_other}", "-", F_3)
    sh.pair("Water-system energy intensity", "w_kwh_m3", "=({pw_mu$}+{pw_treat$})*24/({intake$}*0.00378541)", "kWh/m³", F_3)
    for i, s in enumerate(SCEN):
        cells = ",".join(f'{c}="✅ In POR"' for c in por_cells[s])
        sh.put(f"{'CD'[i]}{sh.r}", f"=IF(AND({cells}),1,0)", "out", F_INT, f"por_ok_{s}")
    sh.put(f"B{sh.r}", "All pumps in POR (1/0)", "label")
    sh.r += 2
    chart = BarChart()
    chart.type, chart.grouping = "col", "clustered"
    chart.title, chart.y_axis.title = "Plant power by component (kW)", "kW"
    chart.add_data(Reference(sh.ws, min_col=3, max_col=4, min_row=pw_first - 1, max_row=pw_first + 3), titles_from_data=True)
    chart.set_categories(Reference(sh.ws, min_col=2, min_row=pw_first, max_row=pw_first + 3))
    chart.height, chart.width = 8, 16
    sh.ws.add_chart(chart, "S5")
    sh.widths({"A": 2, "B": 42, "C": 26, "D": 28, "E": 12, "F": 14, "G": 12, "H": 13, "I": 13,
               "J": 12, "K": 12, "L": 11, "M": 11, "N": 14, "O": 10, "P": 10, "Q": 10})


def build_microbial(book: Book, cfg: Mapping[str, Any]) -> None:
    sh = book.sheet("mic")
    cl, ex, dr = cfg["treatment"]["chlorination"], cfg["microbial_exposure"], cfg["library"]["dose_response"]
    sc = cfg["scenarios"]
    sh.title("Disinfection CT & Quantitative Microbial Risk Assessment",
             "CT = C·T10 · T10 = TDT·BF · LRT = LOG(C_in/C_target) · Pd = 1 − (1 + d/β)^−α · Py = 1 − (1 − Pd)^n")
    sh.section("1. Chlorine contact inputs")
    sh.line("Contact tank operating volume", "cl_vol", float(cl["contact_volume_gal"]), "gal", F_INT)
    sh.line("Peak flow factor", "cl_pf", float(cl["peak_factor"]), "-", F_2)
    sh.line("Baffling factor BF (0.1 unbaffled – 0.9 plug flow)", "cl_bf", float(cl["baffling_factor"]), "-", F_2)
    sh.line("Water temperature", "cl_temp", float(cl["temperature_c"]), "°C", F_1)
    sh.line("Free chlorine residual (Mass Balance outlet)", "cl_res", "={cout_chlorination_cl2}", "mg/L", F_3)

    sh.section("2. CT required lookup (free chlorine, SWTR Guidance Manual)")
    sh.header(["Temp (°C)", "CT 4-log virus", "", "Temp (°C)", "CT 3-log Giardia (pH 7)"])
    r0 = sh.r
    for i in range(max(len(CT_VIRUS_4LOG), len(CT_GIARDIA_3LOG))):
        r = r0 + i
        if i < len(CT_VIRUS_4LOG):
            sh.put(f"B{r}", CT_VIRUS_4LOG[i][0], "in", F_1)
            sh.put(f"C{r}", CT_VIRUS_4LOG[i][1], "in", F_1)
        if i < len(CT_GIARDIA_3LOG):
            sh.put(f"E{r}", CT_GIARDIA_3LOG[i][0], "in", F_1)
            sh.put(f"F{r}", CT_GIARDIA_3LOG[i][1], "in", F_1)
    nv, ng = len(CT_VIRUS_4LOG), len(CT_GIARDIA_3LOG)
    book.name_range("ctv_t", sh.ws, f"B{r0}", f"B{r0 + nv - 1}")
    book.name_range("ctv_v", sh.ws, f"C{r0}", f"C{r0 + nv - 1}")
    book.name_range("ctg_t", sh.ws, f"E{r0}", f"E{r0 + ng - 1}")
    book.name_range("ctg_v", sh.ws, f"F{r0}", f"F{r0 + ng - 1}")
    sh.r = r0 + max(nv, ng) + 1

    sh.section("3. CT verification")
    sh.header(["Parameter", sc["baseline"]["label"], sc["optimized"]["label"], "Unit", "Formula"])
    sh.pair("Contact flow (train)", "cl_q", "={intake$}*{rel_chlorination}", "gpd", F_INT)
    sh.pair("Peak flow Q_peak", "cl_qpk", "={cl_q$}/1440*{cl_pf}", "gpm", F_INT)
    sh.pair("Total detention time TDT = V/Q_peak", "tdt", "={cl_vol}/{cl_qpk$}", "min", F_2)
    sh.pair("T10 = TDT × BF", "t10", "={tdt$}*{cl_bf}", "min", F_2)
    sh.pair("CT = C × T10", "ct", "={cl_res}*{t10$}", "mg·min/L", F_2)
    sh.pair("CT required 4-log virus", "ctv", "=" + x_interp("{cl_temp}", "{ctv_t}", "{ctv_v}", nv), "mg·min/L", F_2)
    sh.pair("CT required 3-log Giardia", "ctg", "=" + x_interp("{cl_temp}", "{ctg_t}", "{ctg_v}", ng), "mg·min/L", F_2)
    sh.pair("Virus inactivation credit = MIN(4, 4·CT/CTreq)", "clv", "={on_chlorination}*MIN(4,4*{ct$}/{ctv$})", "log", F_2)
    sh.pair("Giardia inactivation credit = MIN(3, 3·CT/CTreq)", "clg", "={on_chlorination}*MIN(3,3*{ct$}/{ctg$})", "log", F_2)
    sh.pair("4-log virus achieved?", "ct_chk", "=IF({clv$}>=4,\"✅ Achieved\",\"❌ Not achieved\")", "", kind="check")

    sh.section("4. Physical log-removal credits (enabled nodes, Mass Balance sheet)")
    sh.line("Virus LRV (Σ on × LRV)", "lrvp_virus", "=SUMPRODUCT({on_range},{lrv_virus_range})", "log", F_2)
    sh.line("Protozoa LRV", "lrvp_protozoa", "=SUMPRODUCT({on_range},{lrv_protozoa_range})", "log", F_2)
    sh.line("Bacteria LRV", "lrvp_bacteria", "=SUMPRODUCT({on_range},{lrv_bacteria_range})", "log", F_2)

    sh.section("5. Ingestion QMRA (worker exposure to cooling water)")
    sh.line("Target annual risk Py", "target_py", float(ex["target_annual_risk"]), "per yr", F_SCI)
    sh.line("Ingested volume per event", "ing_v", float(ex["ingestion_volume_l"]), "L", F_SCI, "WERF Table 3-1")
    sh.line("Events per year n", "ing_n", float(ex["ingestion_events_per_year"]), "1/yr", F_INT)
    sh.line("Allowable per-event risk Pd = 1 − (1 − Py)^(1/n)", "pd_allow", "=1-(1-{target_py})^(1/{ing_n})", "-", F_SCI)
    sh.r += 1
    sh.header(["Pathogen", "Model", "α", "β", "r", "LRV class", "C source", "Allowable dose d", "C target",
               "LRT required", "LRV Base", "Py Base", "Check Base", "LRV Opt", "Py Opt", "Check Opt"])
    cls_name = {"virus": "{lrvp_virus}", "protozoa": "{lrvp_protozoa}", "bacteria": "{lrvp_bacteria}"}
    chk_cells: Dict[str, List[str]] = {"b": [], "o": []}
    for p in QMRA_INGESTION:
        prm, r = dr[p], sh.r
        sh.put(f"B{r}", PATHOGEN_LABELS[p], "label")
        sh.put(f"C{r}", prm["model"], "in")
        sh.put(f"D{r}", float(prm.get("alpha", 0.0)), "in", F_3)
        sh.put(f"E{r}", float(prm.get("beta", 0.0)), "in", F_3)
        sh.put(f"F{r}", float(prm.get("r", 0.0)), "in", F_SCI)
        sh.put(f"G{r}", prm["class"], "label")
        sh.put(f"H{r}", f"={{rawp_{p}}}", "out", F_SCI)
        sh.put(f"I{r}", f"=IF(C{r}=\"beta_poisson\",E{r}*((1-{{pd_allow}})^(-1/D{r})-1),-LN(1-{{pd_allow}})/F{r})", "out", F_SCI)
        sh.put(f"J{r}", f"=I{r}/{{ing_v}}", "out", F_SCI)
        sh.put(f"K{r}", f"=IF(H{r}>J{r},LOG(H{r}/J{r}),0)", "out", F_2, f"lrt_{p}")
        credit = {"virus": "{clv$}", "giardia": "{clg$}", "bacteria": "{clv$}", "cryptosporidium": "0"}[CHLORINE_CREDIT_KEY[p]]
        for s, (cl_lrv, cpy, cchk) in zip(SCEN, (("L", "M", "N"), ("O", "P", "Q"))):
            sh.put(f"{cl_lrv}{r}", f"={cls_name[prm['class']]}+{credit.replace('$', '_' + s)}", "out", F_2, f"lrv_{p}_{s}")
            dose = f"(H{r}/10^{cl_lrv}{r}*{{ing_v}})"
            pd = f"IF(C{r}=\"beta_poisson\",1-(1+{dose}/E{r})^(-D{r}),1-EXP(-F{r}*{dose}))"
            sh.put(f"{cpy}{r}", f"=1-(1-{pd})^{{ing_n}}", "out", F_SCI, f"py_{p}_{s}")
            sh.put(f"{cchk}{r}", f"=IF({cpy}{r}<={{target_py}},\"✅\",\"❌\")", "check")
            chk_cells[s].append(f"{cchk}{r}")
        sh.r += 1
    for i, s in enumerate(SCEN):
        cells = ",".join(f'{c}="✅"' for c in chk_cells[s])
        sh.put(f"{'CD'[i]}{sh.r + 1}", f"=IF(AND({cells}),1,0)", "out", F_INT, f"qmra_ok_{s}")
    sh.put(f"B{sh.r + 1}", "All ingestion pathogens meet target (1/0)", "label")
    sh.r += 3

    sh.section("6. Legionella inhalation QMRA — drift aerosol (exponential, r = 0.06)")
    sh.line("Recirculating Legionella (control level)", "leg_c", float(ex["legionella_recirc_cfu_l"]), "CFU/L", F_INT)
    sh.line("Receptor plume dilution", "leg_dil", float(ex["receptor_dilution"]), "-", F_SCI, "Screening")
    sh.line("Respirable fraction of drift water", "leg_resp", float(ex["respirable_fraction"]), "-", F_2, "Screening")
    sh.line("Breathing rate", "leg_br", float(ex["breathing_rate_m3_h"]), "m³/h", F_2)
    sh.line("Exposure duration per event", "leg_hr", float(ex["exposure_hours_per_event"]), "h", F_2)
    sh.line("Events per year", "leg_n", float(ex["inhalation_events_per_year"]), "1/yr", F_INT)
    sh.line("Dose-response r", "leg_r", float(dr["legionella"]["r"]), "-", F_3, "Armstrong & Haas (2007)")
    sh.header(["Parameter", sc["baseline"]["label"], sc["optimized"]["label"], "Unit", "Formula"])
    sh.pair("Aerosol water at receptor", "aer", "={drift_gpd$}/1440*3.78541/60/{airflow$}*{leg_dil}", "L/m³", F_SCI)
    sh.pair("Inhaled water-equivalent volume", "inh_v", "={aer$}*{leg_resp}*{leg_br}*{leg_hr}", "L/event", F_SCI)
    sh.pair("Dose d = C × V", "leg_d", "={leg_c}*{inh_v$}", "CFU", F_SCI)
    sh.pair("Pd = 1 − EXP(−r·d)", "leg_pd", "=1-EXP(-{leg_r}*{leg_d$})", "-", F_SCI)
    sh.pair("Py = 1 − (1 − Pd)^n", "leg_py", "=1-(1-{leg_pd$})^{leg_n}", "per yr", F_SCI)
    sh.pair("Allowable recirculating Legionella", "leg_ct",
            "=-LN(1-(1-(1-{target_py})^(1/{leg_n})))/{leg_r}/{inh_v$}", "CFU/L", F_SCI)
    sh.pair("Required additional LRT on recirc water", "leg_lrt", "=IF({leg_c}>{leg_ct$},LOG({leg_c}/{leg_ct$}),0)", "log", F_2)
    sh.pair("Legionella risk check", "leg_chk", "=IF({leg_py$}<={target_py},\"✅ Py ≤ 1e-4\",\"❌ Py > 1e-4\")", "", kind="check")
    sh.widths({"A": 2, "B": 46, "C": 26, "D": 28, "E": 12, "F": 30, "G": 11, "H": 13, "I": 15, "J": 13,
               "K": 12, "L": 11, "M": 12, "N": 9, "O": 11, "P": 12, "Q": 9})


def build_financial(book: Book, cfg: Mapping[str, Any]) -> None:
    sh = book.sheet("fin")
    eco, ct, sc = cfg["economics"], cfg["cooling_tower"], cfg["scenarios"]
    cx = eco["capex"]
    site = cfg["site"]
    sh.title("Financial & Lifecycle Cost Analysis (NYC DEP FY2022 tariffs)",
             "i = r_interest + r_inflation · PV = FV/(1+i)^n · NPV = Σ CF_t/(1+i)^t − CapEx · SPBT = CapEx / savings")
    sh.section("1. Tariffs, prices & finance inputs")
    sh.line("Water rate", "w_hcf", float(eco["water_rate_per_hcf"]), "$/100 ft³", F_2, "NYC DEP FY22 metered water")
    sh.line("Water rate", "w_kgal", "={w_hcf}/748.052*1000", "$/kgal", F_3)
    sh.line("Sewer charge multiplier", "sew_mult", float(eco["sewer_multiplier"]), "× water", F_2, "159 % of water charge")
    sh.line("Sewer rate", "s_kgal", "={w_kgal}*{sew_mult}", "$/kgal", F_3)
    sh.line("Cooling-tower sewer allowance (evaporation exempt) 1/0", "allow",
            int(bool(eco["cooling_tower_sewer_allowance"])), "-")
    sh.line("Electricity", "elec", float(eco["electricity_per_kwh"]), "$/kWh", F_3)
    sh.line("Sodium hypochlorite (as Cl₂)", "p_naocl", float(eco["naocl_per_kg"]), "$/kg", F_2)
    sh.line("Alum", "p_alum", float(eco["alum_per_kg"]), "$/kg", F_2)
    sh.line("Scale/corrosion inhibitor", "p_inh", float(eco["inhibitor_per_kg"]), "$/kg", F_2)
    sh.line("Sludge disposal", "p_sludge", float(eco["sludge_per_dry_ton"]), "$/dry ton", F_2)
    sh.line("Alum sludge yield", "alum_yield", float(eco["alum_sludge_yield"]), "kg/kg", F_2)
    sh.line("Tower biocide dose (per makeup)", "biocide", float(ct["biocide_dose_mg_l_makeup"]), "mg/L", F_2)
    sh.line("Inhibitor held in recirculation", "inhib", float(ct["inhibitor_mg_l_recirc"]), "mg/L", F_2)
    sh.line("Grid emission factor", "grid", float(eco["grid_lb_co2_per_kwh"]), "lb CO₂/kWh", F_2)
    sh.line("Social cost of carbon", "scc", float(eco["social_cost_carbon_per_tonne"]), "$/t", F_INT)
    sh.line("Hours per year", "hours", float(site["hours_per_year"]), "h", F_INT)
    sh.line("Operating days equivalent = hours/24 × utilization", "days", "={hours}/24*{util}", "d/yr", F_1)
    sh.line("Interest rate", "r_int", float(eco["interest_rate"]), "-", F_PCT)
    sh.line("Inflation rate", "r_inf", float(eco["inflation_rate"]), "-", F_PCT)
    sh.line("Discount rate i = interest + inflation", "i_rate", "={r_int}+{r_inf}", "-", F_PCT)
    sh.line("Utility cost escalation", "esc", float(eco["escalation_rate"]), "-", F_PCT, "NYC water rate +2.76 % (FY22)")
    sh.line("Fill cost", "cx_fill", float(cx["fill_per_m3"]), "$/m³", F_INT)
    sh.line("Drift eliminator cost", "cx_de", float(cx["eliminator_per_m2"]), "$/m²", F_INT)
    sh.line("Blowdown-recovery RO cost", "cx_bdro", float(cx["blowdown_ro_per_gpd"]), "$/gpd", F_2)
    sh.line("Installation factor", "cx_inst", float(cx["installation_factor"]), "×", F_2)

    sh.section("2. Annual quantities")
    sh.header(["Quantity", sc["baseline"]["label"], sc["optimized"]["label"], "Unit", ""])
    sh.pair("Raw water intake", "a_intake", "={intake$}*{days}", "gal/yr", F_INT)
    sh.pair("Sewer discharge", "a_sewer", "={sewer_gpd$}*{days}", "gal/yr", F_INT)
    sh.pair("Electricity", "a_kwh", "={kw_total$}*{hours}*{util}", "kWh/yr", F_INT)
    sh.pair("Alum", "a_alum", "={on_coag_sed}*{alum_dose}*{intake$}*{rel_coag_sed}*3.78541/1E6*{days}", "kg/yr", F_INT)
    sh.pair("NaOCl (contact + tower biocide)", "a_naocl",
            "=({on_chlorination}*{D_chlorination_cl2}*{intake$}*{rel_chlorination}+{biocide}*{mt$})*3.78541/1E6*{days}",
            "kg/yr", F_INT)
    sh.pair("Inhibitor = C × (B + D)", "a_inh", "={inhib}*({bd$}+{drift_gpd$})*3.78541/1E6*{days}", "kg/yr", F_INT)
    sh.pair("Dry sludge", "a_sludge",
            "=(SUM({tss_removed_range})*{intake$}+{a_alum$}/{days}*{alum_yield})/907.185*{days}", "dry t/yr", F_1)
    sh.pair("CO₂e", "co2", "={a_kwh$}*{grid}/2204.6", "t/yr", F_INT)

    sh.section("3. Annual operating cost (year 1)")
    sh.header(["Cost item", sc["baseline"]["label"], sc["optimized"]["label"], "Unit", ""])
    sh.pair("Water purchase", "cost_water", "=IF({src_purchased}=1,{a_intake$}/1000*{w_kgal},0)", "$/yr", F_USD)
    sh.pair("Sewer", "cost_sewer", "=IF({allow}=1,{a_sewer$},{a_intake$})/1000*{s_kgal}", "$/yr", F_USD)
    sh.pair("Electricity", "cost_elec", "={a_kwh$}*{elec}", "$/yr", F_USD)
    sh.pair("Chemicals", "cost_chem", "={a_alum$}*{p_alum}+{a_naocl$}*{p_naocl}+{a_inh$}*{p_inh}", "$/yr", F_USD)
    sh.pair("Sludge disposal", "cost_sludge", "={a_sludge$}*{p_sludge}", "$/yr", F_USD)
    sh.pair("Carbon (social cost)", "cost_carbon", "={co2$}*{scc}", "$/yr", F_USD)
    sh.pair("TOTAL", "cost_total", "=SUM({cost_water$},{cost_sewer$},{cost_elec$},{cost_chem$},{cost_sludge$},{cost_carbon$})",
            "$/yr", F_USD)
    sh.line("Year-1 net operating savings (Baseline − Optimized)", "savings", "={cost_total_b}-{cost_total_o}", "$/yr", F_USD)

    sh.section("4. Upgrade capital cost")
    sh.line("Corrugated fill = area × height × $/m³", "capex_fill", "={area_o}*{fh_o}*{cx_fill}", "$", F_USD)
    sh.line("HE drift eliminators = area × $/m²", "capex_de", "={area_o}*{cx_de}", "$", F_USD)
    sh.line("Blowdown-recovery RO", "capex_bdro", "=IF(AND({bdrec_o}=1,{bdrec_b}=0),{bd_o}*{cx_bdro},0)", "$", F_USD)
    sh.line("Installation", "capex_inst", "=({capex_fill}+{capex_de}+{capex_bdro})*({cx_inst}-1)", "$", F_USD)
    sh.line("TOTAL CapEx", "capex_total", "={capex_fill}+{capex_de}+{capex_bdro}+{capex_inst}", "$", F_USD)

    years = int(eco["horizon_years"])
    sh.section(f"5. {years}-year discounted cash flow", span=7)
    sh.header(["Year t", "Cash flow CF_t", "Discount factor 1/(1+i)^t", "PV = CF/(1+i)^t", "Cumulative PV"])
    r0 = sh.r
    sh.put(f"B{r0}", 0, "label")
    sh.put(f"C{r0}", "=-{capex_total}", "out", F_USD)
    sh.put(f"D{r0}", 1, "out", F_4)
    sh.put(f"E{r0}", f"=C{r0}*D{r0}", "out", F_USD)
    sh.put(f"F{r0}", f"=E{r0}", "out", F_USD)
    for t in range(1, years + 1):
        r = r0 + t
        sh.put(f"B{r}", t, "label")
        sh.put(f"C{r}", f"={{savings}}*(1+{{esc}})^(B{r}-1)", "out", F_USD)
        sh.put(f"D{r}", f"=1/(1+{{i_rate}})^B{r}", "out", F_4)
        sh.put(f"E{r}", f"=C{r}*D{r}", "out", F_USD)
        sh.put(f"F{r}", f"=F{r - 1}+E{r}", "out", F_USD)
    rl = r0 + years
    sh.r = rl + 2
    sh.line("NPV = −CapEx + NPV(i, CF₁…CF_N)", "npv", f"=C{r0}+NPV({{i_rate}},C{r0 + 1}:C{rl})", "$", F_USD)
    sh.line("NPV check = Σ PV", None, f"=SUM(E{r0}:E{rl})", "$", F_USD)
    sh.line("IRR", "irr", f"=IFERROR(IRR(C{r0}:C{rl}),\"n/a\")", "-", F_PCT)
    k = f"COUNTIF(F{r0 + 1}:F{rl},\"<0\")"
    sh.line("Discounted break-even year (BEP)", "bep",
            f"=IF({k}={years},\"> horizon\",{k}+(-INDEX(F{r0}:F{rl},{k}+1))/(INDEX(F{r0}:F{rl},{k}+2)-INDEX(F{r0}:F{rl},{k}+1)))",
            "years", F_2)
    sh.line("Simple payback SPBT = CapEx / savings", "spbt", "=IF({savings}>0,{capex_total}/{savings},\"n/a\")", "years", F_2)
    sh.line("Investment decision", None, "=IF({npv}>0,\"✅ Positive NPV — invest\",\"❌ Negative NPV\")", "", kind="check")

    line = LineChart()
    line.title, line.y_axis.title, line.x_axis.title = "Cumulative discounted cash flow", "$", "Year"
    line.add_data(Reference(sh.ws, min_col=6, min_row=r0 - 1, max_row=rl), titles_from_data=True)
    line.set_categories(Reference(sh.ws, min_col=2, min_row=r0, max_row=rl))
    line.height, line.width = 8, 16
    sh.ws.add_chart(line, "H5")
    bar = BarChart()
    bar.type, bar.title = "col", "Annual savings by year"
    bar.add_data(Reference(sh.ws, min_col=3, min_row=r0, max_row=rl), titles_from_data=True)
    bar.set_categories(Reference(sh.ws, min_col=2, min_row=r0 + 1, max_row=rl))
    bar.height, bar.width = 8, 16
    sh.ws.add_chart(bar, "H22")
    sh.widths({"A": 2, "B": 50, "C": 26, "D": 28, "E": 18, "F": 18})


# =========================================================================================
def build_workbook(cfg: Mapping[str, Any]) -> Workbook:
    """Create the full formula-driven workbook from a config dict."""
    return build_book(cfg).wb


def build_book(cfg: Mapping[str, Any]) -> Book:
    """Build the workbook and return it with its named-reference registry."""
    book = Book()
    dash = build_dashboard(book, cfg)
    build_mass_balance(book, cfg)
    build_cooling(book, cfg)
    build_hydraulics(book, cfg)
    build_microbial(book, cfg)
    build_financial(book, cfg)
    build_source_library(book, cfg)
    build_merkel(book)
    finish_dashboard(dash)
    add_dashboard_validation(book, dash, cfg)
    book.finalize()
    book.wb.calculation = CalcProperties(fullCalcOnLoad=True)
    for ws in book.wb.worksheets:
        ws.sheet_properties.tabColor = "1F3864" if ws.title == SHEETS["dash"] else (
            "BF9000" if ws.title in (SHEETS["src"], SHEETS["mk"]) else "548235")
    return book


def export_workbook(cfg: Mapping[str, Any], path: Union[str, Path] = DEFAULT_FILENAME) -> Path:
    out = Path(path)
    build_workbook(cfg).save(out)
    return out


def workbook_bytes(cfg: Mapping[str, Any]) -> bytes:
    """Workbook as bytes (for the Streamlit download button)."""
    buf = io.BytesIO()
    build_workbook(cfg).save(buf)
    return buf.getvalue()
